"""
SAMC — OCR 분석 결과를 DOCX / PDF로 내보내는 서비스.

포함 섹션:
  1. 기본 정보
  2. 원재료 배합비율 (성분코드 포함)
  3. 제조공정 — 단계별 (추천코드 + 유사코드) 표시
  4. 수출국 라벨 분석 (디자인 설명, 라벨 문구, 경고)
  5. 라벨 제품 이미지 (Vision 추출 이미지 + 텍스트 필드)

- DOCX: python-docx (한국어 완벽 지원)
- PDF: fpdf2 (팀 컨벤션, 한국어 TTF 자동 탐지)
"""

from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import Any

logger = logging.getLogger(__name__)

# 라벨 이미지 텍스트 필드 레이블
_IMG_FIELD_LABELS = [
    ("label_product_name",   "제품명"),
    ("label_ingredients",    "원재료"),
    ("label_content_volume", "내용량"),
    ("label_origin",         "원산지"),
    ("label_manufacturer",   "제조사"),
    ("label_case_number",    "케이스 넘버"),
]


# ─────────────────────────────────────────────
# DOCX 생성
# ─────────────────────────────────────────────

def build_docx(
    parsed: dict[str, Any],
    *,
    product_name: str = "",
    case_id: str = "",
    label_images: list[dict] | None = None,
) -> bytes:
    from docx import Document
    from docx.shared import Pt, Cm, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    label_images = label_images or []
    doc = Document()

    # 기본 폰트 (한국어 안전)
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10)

    basic = parsed.get("basic_info", {}) or {}
    pname = basic.get("product_name") or product_name or ""
    ings  = parsed.get("ingredients", []) or []
    proc  = parsed.get("process_info", {}) or {}
    label = parsed.get("label_info", {}) or {}

    # ── 표지 헤더 ──
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run("SAMC 수입식품 OCR 분석 결과")
    run.bold = True
    run.font.size = Pt(18)

    if pname:
        sub_p = doc.add_paragraph()
        sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub_p.add_run(pname)
        sub_run.font.size = Pt(13)
        sub_run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.add_run(
        f"생성일: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
        + (f"   |   Case ID: {case_id}" if case_id else "")
    ).font.size = Pt(9)

    doc.add_paragraph()

    # ── 1. 기본 정보 ──
    doc.add_heading("1. 기본 정보", level=1)
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in [
        ("제품명",        basic.get("product_name") or product_name or "-"),
        ("수출국",        basic.get("export_country") or "-"),
        ("최초 수입 여부", "예" if basic.get("is_first_import") else "아니오"),
        ("유기인증",      "예" if basic.get("is_organic") else "아니오"),
        ("OEM",          "예" if basic.get("is_oem") else "아니오"),
    ]:
        row = t.add_row().cells
        row[0].text = k
        row[1].text = str(v)
    doc.add_paragraph()

    # ── 2. 원재료 배합비율 ──
    doc.add_heading("2. 원재료 배합비율", level=1)
    if ings:
        tbl = doc.add_table(rows=1, cols=7)
        tbl.style = "Light Grid Accent 1"
        hdr = tbl.rows[0].cells
        for i, h in enumerate(["성분명", "비율(%)", "원산지", "INS", "CAS", "식약처코드", "공식성분명"]):
            hdr[i].text = h
            for p in hdr[i].paragraphs:
                for r in p.runs:
                    r.bold = True
        for ing in ings:
            r = tbl.add_row().cells
            r[0].text = str(ing.get("name") or "")
            r[1].text = str(ing.get("ratio") or "")
            r[2].text = str(ing.get("origin") or "")
            r[3].text = str(ing.get("ins_number") or "")
            r[4].text = str(ing.get("cas_number") or "")
            r[5].text = str(ing.get("ingredient_code") or "")
            r[6].text = str(ing.get("ingredient_code_name") or "")
    else:
        doc.add_paragraph("(추출된 원재료 없음)")
    doc.add_paragraph()

    # ── 3. 제조공정 ──
    doc.add_heading("3. 제조공정", level=1)
    codes = proc.get("process_codes") or []
    steps = proc.get("process_steps") or []
    raw   = proc.get("raw_process_text") or ""
    is_incomplete = proc.get("is_incomplete", False)

    p = doc.add_paragraph()
    p.add_run("전체 공정 코드: ").bold = True
    p.add_run(", ".join(codes) if codes else "(없음)")

    if is_incomplete:
        warn_p = doc.add_paragraph()
        warn_p.add_run("⚠ 파싱 불완전: ").bold = True
        warn_p.add_run(str(proc.get("incomplete_reason") or "일부 공정 코드만 추출됨. 직접 입력/수정 필요."))

    # 단계별 공정 코드 표 (추천코드 + 유사코드)
    if steps:
        doc.add_paragraph("단계별 공정 분석:").runs[0].bold = True
        for step in steps:
            snum  = step.get("step_number", "")
            sorig = step.get("step_name_original", "")
            sko   = step.get("step_name_ko", "")
            rcode = step.get("recommended_code", "")
            rname = step.get("recommended_code_name", "")
            rrsn  = step.get("recommended_reason", "")
            sims  = step.get("similar_codes") or []

            step_p = doc.add_paragraph()
            step_p.add_run(f"{snum}단계: {sorig}").bold = True
            if sko and sko != sorig:
                step_p.add_run(f"  ({sko})")

            rec_p = doc.add_paragraph(style="List Bullet")
            rec_p.add_run(f"추천: {rcode} - {rname}").bold = True
            if rrsn:
                rec_p.add_run(f"  |  {rrsn}")

            for sc in sims:
                sim_p = doc.add_paragraph(style="List Bullet 2")
                sc_code = sc.get("code", "")
                sc_name = sc.get("name", "")
                sc_note = sc.get("confusion_note", "") or sc.get("reason", "")
                sim_p.add_run(f"유사: {sc_code} - {sc_name}")
                if sc_note:
                    sim_p.add_run(f"  →  {sc_note}").font.color.rgb = __import__("docx.shared", fromlist=["RGBColor"]).RGBColor(0x64, 0x74, 0x8b)

    if raw:
        doc.add_paragraph("OCR 추출 원문:").runs[0].bold = True
        for line in raw.splitlines():
            doc.add_paragraph(line or " ")
    doc.add_paragraph()

    # ── 4. 수출국 라벨 분석 ──
    doc.add_heading("4. 수출국 라벨 분석", level=1)
    ltexts = label.get("label_texts") or []
    warns  = label.get("warnings") or []
    desc   = label.get("design_description") or ""

    if desc:
        doc.add_paragraph("디자인 설명:").runs[0].bold = True
        doc.add_paragraph(desc)

    if ltexts:
        doc.add_paragraph("라벨 문구:").runs[0].bold = True
        for lt in ltexts:
            doc.add_paragraph(str(lt), style="List Bullet")

    if warns:
        doc.add_paragraph("경고/주의사항:").runs[0].bold = True
        for w in warns:
            doc.add_paragraph(str(w), style="List Bullet")

    if not (ltexts or warns or desc):
        doc.add_paragraph("(라벨 정보 없음)")
    doc.add_paragraph()

    # ── 5. 라벨 제품 이미지 ──
    if label_images:
        doc.add_heading("5. 라벨 제품 이미지", level=1)
        doc.add_paragraph(
            f"Vision AI가 자동 추출한 제품 이미지 {len(label_images)}개"
        ).runs[0].font.size = Pt(9)

        for img_data in label_images:
            img_bytes = img_data.get("bytes")
            idx = img_data.get("image_index", 0)

            doc.add_heading(f"이미지 {idx + 1}", level=2)

            # 이미지 삽입
            if img_bytes:
                try:
                    doc.add_picture(io.BytesIO(img_bytes), width=Cm(8))
                except Exception as e:
                    logger.warning(f"DOCX 이미지 삽입 실패 (idx={idx}): {e}")
                    doc.add_paragraph("(이미지 삽입 실패)")

            # 추출 텍스트 필드 테이블
            fields = [(lbl, img_data.get(key)) for key, lbl in _IMG_FIELD_LABELS if img_data.get(key)]
            if fields:
                ft = doc.add_table(rows=0, cols=2)
                ft.style = "Light Grid Accent 1"
                for lbl, val in fields:
                    row = ft.add_row().cells
                    row[0].text = lbl
                    row[0].paragraphs[0].runs[0].bold = True
                    row[1].text = str(val)
            else:
                doc.add_paragraph("(추출된 텍스트 없음)").runs[0].font.size = Pt(9)

            doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────
# PDF 생성 (fpdf2 — 팀 컨벤션)
# ─────────────────────────────────────────────

def _find_korean_font() -> str | None:
    """시스템에서 한국어 지원 TTF 폰트 경로를 탐지해 반환."""
    import os
    candidates = [
        # Windows (맑은 고딕)
        os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts", "malgun.ttf"),
        os.path.join(os.environ.get("WINDIR", "C:\\Windows"), "Fonts", "gulim.ttc"),
        # Linux / Ubuntu (나눔고딕)
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
        "/usr/share/fonts/truetype/nanum/NanumBarunGothic.ttf",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        # macOS
        "/Library/Fonts/AppleGothic.ttf",
        "/System/Library/Fonts/AppleSDGothicNeo.ttc",
        # 프로젝트 번들 (backend/fonts/ 에 추가 시)
        os.path.normpath(os.path.join(os.path.dirname(__file__), "..", "fonts", "NanumGothic.ttf")),
    ]
    for p in candidates:
        try:
            if os.path.isfile(os.path.normpath(p)):
                return os.path.normpath(p)
        except Exception:
            pass
    return None


def build_pdf(
    parsed: dict[str, Any],
    *,
    product_name: str = "",
    case_id: str = "",
    label_images: list[dict] | None = None,
) -> bytes:
    from fpdf import FPDF

    label_images = label_images or []

    # ── 한국어 폰트 탐지 ──
    font_path = _find_korean_font()
    if not font_path:
        logger.warning("한국어 TTF 폰트를 찾지 못했습니다. backend/fonts/NanumGothic.ttf 를 추가하거나 시스템 폰트를 설치하세요.")

    # ── fpdf2 PDF 초기화 ──
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(left=18, top=18, right=18)

    # 폰트 등록
    KR = "KR"
    if font_path:
        try:
            pdf.add_font(KR, fname=font_path)
            _font_ok = True
        except Exception as e:
            logger.warning(f"한국어 폰트 등록 실패: {e}")
            KR = "Helvetica"
            _font_ok = False
    else:
        KR = "Helvetica"
        _font_ok = False

    def _set(size: int, bold: bool = False):
        pdf.set_font(KR, size=size)
        if bold:
            pdf.set_text_color(15, 23, 42)
        else:
            pdf.set_text_color(30, 30, 30)

    def _row(label: str, value: str, lw: float = 40, vw: float = 130):
        """키-값 한 행 출력."""
        _set(9, bold=True)
        pdf.set_fill_color(241, 245, 249)
        pdf.cell(lw, 7, label, border=1, fill=True)
        _set(9)
        pdf.set_fill_color(255, 255, 255)
        pdf.multi_cell(vw, 7, value, border=1)

    def _section(title: str):
        pdf.ln(4)
        _set(13, bold=True)
        pdf.set_fill_color(15, 23, 42)
        pdf.set_text_color(255, 255, 255)
        pdf.cell(0, 9, f"  {title}", border=0, ln=True, fill=True)
        pdf.set_text_color(30, 30, 30)
        pdf.ln(2)

    def _bullet(text: str, indent: float = 6, size: int = 9):
        _set(size)
        pdf.set_x(pdf.get_x() + indent)
        pdf.multi_cell(0, 6, f"• {text}", border=0)

    def _text(text: str, size: int = 9):
        _set(size)
        pdf.multi_cell(0, 6, text, border=0)

    basic = parsed.get("basic_info", {}) or {}
    pname = basic.get("product_name") or product_name or ""
    ings  = parsed.get("ingredients", []) or []
    proc  = parsed.get("process_info", {}) or {}
    label = parsed.get("label_info", {}) or {}

    pdf.add_page()

    # ── 표지 헤더 ──
    _set(18, bold=True)
    pdf.set_text_color(15, 23, 42)
    pdf.cell(0, 12, "SAMC 수입식품 OCR 분석 결과", ln=True, align="C")
    if pname:
        _set(13)
        pdf.cell(0, 8, pname, ln=True, align="C")
    _set(9)
    pdf.set_text_color(100, 116, 139)
    meta_txt = f"생성일: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    if case_id:
        meta_txt += f"   |   Case ID: {case_id}"
    pdf.cell(0, 6, meta_txt, ln=True, align="C")
    pdf.set_text_color(30, 30, 30)
    pdf.ln(6)

    # ── 1. 기본 정보 ──
    _section("1. 기본 정보")
    for k, v in [
        ("제품명",        basic.get("product_name") or product_name or "-"),
        ("수출국",        basic.get("export_country") or "-"),
        ("제조사",        basic.get("manufacturer") or "-"),
        ("내용량",        basic.get("content_volume") or "-"),
        ("최초 수입",     "예" if basic.get("is_first_import") else "아니오"),
        ("유기인증",      "예" if basic.get("is_organic") else "아니오"),
        ("OEM",          "예" if basic.get("is_oem") else "아니오"),
    ]:
        _row(k, str(v))

    # ── 2. 원재료 배합비율 ──
    _section("2. 원재료 배합비율")
    if ings:
        _set(8, bold=True)
        pdf.set_fill_color(15, 23, 42)
        pdf.set_text_color(255, 255, 255)
        col_ws = [45, 16, 22, 16, 28, 24, 29]
        for h_txt, cw in zip(["성분명", "비율%", "원산지", "INS", "CAS", "식약처코드", "공식성분명"], col_ws):
            pdf.cell(cw, 7, h_txt, border=1, fill=True)
        pdf.ln()
        pdf.set_text_color(30, 30, 30)
        _set(8)
        for i, ing in enumerate(ings):
            fill_color = (248, 250, 252) if i % 2 == 1 else (255, 255, 255)
            pdf.set_fill_color(*fill_color)
            row_data = [
                str(ing.get("name") or ""),
                str(ing.get("ratio") or ""),
                str(ing.get("origin") or ""),
                str(ing.get("ins_number") or ""),
                str(ing.get("cas_number") or ""),
                str(ing.get("ingredient_code") or ""),
                str(ing.get("ingredient_code_name") or ""),
            ]
            for val, cw in zip(row_data, col_ws):
                pdf.cell(cw, 7, val[:20], border=1, fill=True)
            pdf.ln()
    else:
        _text("(추출된 원재료 없음)")

    # ── 3. 제조공정 ──
    _section("3. 제조공정")
    codes = proc.get("process_codes") or []
    steps = proc.get("process_steps") or []
    raw   = proc.get("raw_process_text") or ""
    is_incomplete = proc.get("is_incomplete", False)

    _set(9, bold=True)
    pdf.multi_cell(0, 6, f"전체 공정 코드: {', '.join(codes) if codes else '(없음)'}")

    if is_incomplete:
        pdf.ln(2)
        _set(9)
        pdf.set_text_color(220, 38, 38)
        pdf.multi_cell(0, 6, f"⚠ 파싱 불완전: {proc.get('incomplete_reason') or '일부만 추출. 직접 입력 필요.'}")
        pdf.set_text_color(30, 30, 30)

    if steps:
        pdf.ln(3)
        _set(9, bold=True)
        pdf.multi_cell(0, 6, "단계별 공정 분석:")
        for step in steps:
            snum  = step.get("step_number", "")
            sorig = step.get("step_name_original", "")
            sko   = step.get("step_name_ko", "")
            rcode = step.get("recommended_code", "")
            rname = step.get("recommended_code_name", "")
            rrsn  = step.get("recommended_reason", "")
            sims  = step.get("similar_codes") or []

            pdf.ln(2)
            _set(9, bold=True)
            label_txt = f"{snum}단계: {sorig}"
            if sko and sko != sorig:
                label_txt += f"  ({sko})"
            pdf.multi_cell(0, 6, label_txt)

            _set(9)
            pdf.set_x(pdf.get_x() + 8)
            pdf.set_text_color(5, 150, 105)   # 초록
            pdf.multi_cell(0, 6, f"  추천: {rcode} - {rname}   |   {rrsn}")
            pdf.set_text_color(30, 30, 30)

            for sc in sims:
                sc_code = sc.get("code", "")
                sc_name = sc.get("name", "")
                sc_note = sc.get("confusion_note", "") or sc.get("reason", "")
                pdf.set_x(pdf.get_x() + 14)
                pdf.set_text_color(100, 116, 139)  # 회색
                note_txt = f"  유사: {sc_code} - {sc_name}"
                if sc_note:
                    note_txt += f"   →   {sc_note}"
                pdf.multi_cell(0, 5, note_txt)
                pdf.set_text_color(30, 30, 30)

    # ── 4. 수출국 라벨 분석 ──
    _section("4. 수출국 라벨 분석")
    ltexts = label.get("label_texts") or []
    warns  = label.get("warnings") or []
    desc   = label.get("design_description") or ""

    if not (ltexts or warns or desc):
        _text("(라벨 정보 없음)")
    else:
        if desc:
            _set(9, bold=True); pdf.multi_cell(0, 6, "디자인 설명:")
            _text(desc)
        if ltexts:
            _set(9, bold=True); pdf.multi_cell(0, 6, "라벨 문구:")
            for lt in ltexts:
                _bullet(str(lt))
        if warns:
            _set(9, bold=True); pdf.multi_cell(0, 6, "경고/주의사항:")
            for w in warns:
                _bullet(str(w))

    # ── 5. 라벨 제품 이미지 ──
    if label_images:
        _section(f"5. 라벨 제품 이미지 ({len(label_images)}개)")
        for img_data in label_images:
            img_bytes = img_data.get("bytes")
            idx = img_data.get("image_index", 0)
            _set(10, bold=True)
            pdf.multi_cell(0, 7, f"이미지 {idx + 1}")
            if img_bytes:
                try:
                    tmp = io.BytesIO(img_bytes)
                    pdf.image(tmp, w=70)
                except Exception as e:
                    logger.warning(f"PDF 이미지 삽입 실패 (idx={idx}): {e}")
                    _text("(이미지 삽입 실패)")
            fields = [(lbl, img_data.get(key)) for key, lbl in _IMG_FIELD_LABELS if img_data.get(key)]
            for lbl, val in fields:
                _row(lbl, str(val), lw=30, vw=140)
            pdf.ln(4)

    return pdf.output()
