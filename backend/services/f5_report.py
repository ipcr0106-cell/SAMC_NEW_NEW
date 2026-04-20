"""
F5 한글표시사항 검토내역서 리포트 생성.

제공 형식:
  - DOCX : python-docx 사용
  - PDF  : reportlab 사용 (fpdf2 대신 - 한글 처리가 훨씬 안정적)

산출물 구조:
  1. 기본 정보 (제품명/확정자/확정일)
  2. 교차검증 종합 (요약 통계)
  3. 항목별 교차검증 결과
  4. AI 추가 발견 이슈 (있는 경우)
  5. 최종 한글표시사항 시안
"""

from __future__ import annotations

import io
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

# reportlab imports
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.colors import HexColor, Color
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    KeepTogether,
    PageBreak,
)


# ════════════════════════════════════════════════════════════
# 공통 헬퍼
# ════════════════════════════════════════════════════════════

DRAFT_LABELS: Dict[str, str] = {
    "product_name":      "제품명",
    "food_type":         "식품유형",
    "ingredients":       "원재료명 및 함량",
    "net_weight":        "내용량",
    "expiry":            "소비기한",
    "storage":           "보관방법",
    "manufacturer":      "제조사",
    "importer":          "수입자",
    "allergy":           "알레르기",
    "gmo":               "GMO",
    "country_of_origin": "원산지",
}

STATUS_LABELS: Dict[str, str] = {
    "pass":    "적합",
    "fail":    "부적합",
    "unclear": "확인필요",
}

CROSS_LABELS: Dict[str, str] = {
    "agree":            "일치",
    "disagree":         "불일치",
    "additional_issue": "추가이슈",
}

SEVERITY_LABELS: Dict[str, str] = {
    "error":   "오류",
    "warning": "경고",
    "info":    "정보",
}


def to_safe_string(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, (int, float, bool)):
        return str(value)
    if isinstance(value, list):
        return ", ".join(to_safe_string(v) for v in value)
    if isinstance(value, dict):
        return "; ".join(f"{k}: {to_safe_string(v)}" for k, v in value.items())
    return str(value)


def safe_filename(text: str) -> str:
    text = (text or "").strip()
    text = re.sub(r"[^0-9A-Za-z가-힣\-\.]+", "_", text)
    text = re.sub(r"_+", "_", text).strip("_")
    return text or "untitled"


def make_filename(
    product_name: str,
    confirmed_at: Optional[str],
    ext: str,
) -> str:
    date_part = ""
    if confirmed_at:
        try:
            dt = datetime.fromisoformat(confirmed_at.replace("Z", "+00:00"))
            date_part = dt.strftime("%Y%m%d")
        except Exception:
            pass
    if not date_part:
        date_part = datetime.now().strftime("%Y%m%d")

    safe_name = safe_filename(product_name or "검토내역서")
    return f"한글표시사항_{safe_name}_{date_part}.{ext}"


def extract_report_data(step_record: Dict[str, Any]) -> Dict[str, Any]:
    final_result: Dict[str, Any] = step_record.get("final_result") or {}
    phase1: Dict[str, Any] = final_result.get("phase1") or {}
    phase2: Dict[str, Any] = final_result.get("phase2") or {}
    draft: Dict[str, Any] = phase2.get("draft") or {}

    items: List[Dict[str, Any]] = phase1.get("items") or []
    validation: List[Dict[str, Any]] = phase2.get("validation") or []
    additional_issues: List[Dict[str, Any]] = phase2.get("additional_issues") or []

    v2_by_field = {v.get("field"): v for v in validation}

    fail_count = sum(1 for i in items if i.get("status") == "fail")
    unclear_count = sum(1 for i in items if i.get("status") == "unclear")
    disagree_count = sum(1 for v in validation if v.get("cross_result") == "disagree")
    error_count = sum(1 for a in additional_issues if a.get("severity") == "error")

    product_name = to_safe_string(draft.get("product_name")) or "(제품명 미기재)"

    return {
        "product_name": product_name,
        "confirmed_by": to_safe_string(final_result.get("confirmed_by")),
        "updated_at": step_record.get("updated_at") or step_record.get("created_at"),
        "items": items,
        "v2_by_field": v2_by_field,
        "additional_issues": additional_issues,
        "draft": draft,
        "stats": {
            "total": len(items),
            "fail": fail_count,
            "unclear": unclear_count,
            "disagree": disagree_count,
            "error": error_count,
        },
    }


# ════════════════════════════════════════════════════════════
# DOCX 생성 (기존 그대로)
# ════════════════════════════════════════════════════════════

def _docx_set_font(run, size: int = 10, bold: bool = False, color: Optional[Tuple[int, int, int]] = None):
    run.font.name = "맑은 고딕"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rFonts is None:
        from docx.oxml.ns import qn
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    from docx.oxml.ns import qn
    rFonts.set(qn("w:eastAsia"), "맑은 고딕")


def _docx_add_heading(doc: Document, text: str, size: int = 13):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"■ {text}")
    _docx_set_font(run, size=size, bold=True, color=(30, 64, 120))


def _docx_add_text(doc: Document, text: str, size: int = 10, bold: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _docx_set_font(run, size=size, bold=bold)


def generate_docx(step_record: Dict[str, Any]) -> Tuple[bytes, str]:
    data = extract_report_data(step_record)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trun = title.add_run("한글표시사항 검토내역서")
    _docx_set_font(trun, size=18, bold=True, color=(20, 20, 20))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = subtitle.add_run("SAMC 수입식품 검역 AI 플랫폼")
    _docx_set_font(srun, size=10, color=(120, 120, 120))

    doc.add_paragraph()

    _docx_add_heading(doc, "기본 정보")
    info_table = doc.add_table(rows=3, cols=2)
    info_table.style = "Table Grid"
    info_rows = [
        ("제품명", data["product_name"]),
        ("확정자", data["confirmed_by"] or "(미확정)"),
        ("확정일시", str(data["updated_at"] or "-")),
    ]
    for i, (label, value) in enumerate(info_rows):
        c0, c1 = info_table.rows[i].cells
        c0.text = ""
        c1.text = ""
        _docx_set_font(c0.paragraphs[0].add_run(label), size=10, bold=True)
        _docx_set_font(c1.paragraphs[0].add_run(str(value)), size=10)

    _docx_add_heading(doc, "교차검증 종합")
    stats = data["stats"]
    summary_line = (
        f"검토항목 {stats['total']}개 / "
        f"법령 부적합 {stats['fail']}건 / "
        f"확인필요 {stats['unclear']}건 / "
        f"1·2차 불일치 {stats['disagree']}건 / "
        f"추가 오류 {stats['error']}건"
    )
    _docx_add_text(doc, summary_line, size=10)

    _docx_add_heading(doc, "항목별 교차검증 결과")
    if data["items"]:
        tbl = doc.add_table(rows=1, cols=5)
        tbl.style = "Table Grid"
        hdr_cells = tbl.rows[0].cells
        headers = ["항목", "1차 판정", "AI 교차", "검토 의견", "법령 근거"]
        for idx, h in enumerate(headers):
            hdr_cells[idx].text = ""
            _docx_set_font(hdr_cells[idx].paragraphs[0].add_run(h), size=10, bold=True)

        for item in data["items"]:
            row = tbl.add_row().cells
            field = to_safe_string(item.get("field"))
            status = item.get("status", "")
            status_kr = STATUS_LABELS.get(status, status)

            v2 = data["v2_by_field"].get(item.get("field")) or {}
            cross = v2.get("cross_result", "")
            cross_kr = CROSS_LABELS.get(cross, cross or "-")

            note = to_safe_string(item.get("note"))
            law_ref = to_safe_string(item.get("law_ref"))

            values = [field, status_kr, cross_kr, note, law_ref]
            for idx, v in enumerate(values):
                row[idx].text = ""
                _docx_set_font(row[idx].paragraphs[0].add_run(v), size=9)
    else:
        _docx_add_text(doc, "(검토 항목 없음)", size=10)

    if data["additional_issues"]:
        _docx_add_heading(doc, "AI 추가 발견 이슈")
        for issue in data["additional_issues"]:
            sev = issue.get("severity", "info")
            sev_kr = SEVERITY_LABELS.get(sev, sev)
            field = to_safe_string(issue.get("field"))
            text = to_safe_string(issue.get("issue"))
            _docx_add_text(doc, f"[{sev_kr}] {field} - {text}", size=10)

    _docx_add_heading(doc, "최종 한글표시사항 시안")
    if data["draft"]:
        draft_tbl = doc.add_table(rows=0, cols=2)
        draft_tbl.style = "Table Grid"
        for key, val in data["draft"].items():
            label = DRAFT_LABELS.get(key, key)
            row = draft_tbl.add_row().cells
            row[0].text = ""
            row[1].text = ""
            _docx_set_font(row[0].paragraphs[0].add_run(label), size=10, bold=True)
            _docx_set_font(row[1].paragraphs[0].add_run(to_safe_string(val)), size=10)
    else:
        _docx_add_text(doc, "(시안 데이터 없음)", size=10)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    filename = make_filename(data["product_name"], data["updated_at"], "docx")
    return buf.getvalue(), filename


# ════════════════════════════════════════════════════════════
# PDF 생성 (reportlab 기반)
# ════════════════════════════════════════════════════════════

_KOREAN_FONT_CANDIDATES: List[Path] = [
    # Windows
    Path("C:/Windows/Fonts/malgun.ttf"),
    Path("C:/Windows/Fonts/malgunbd.ttf"),
    Path("C:/Windows/Fonts/NanumGothic.ttf"),
    # macOS
    Path("/Library/Fonts/AppleSDGothicNeo.ttc"),
    Path("/System/Library/Fonts/AppleSDGothicNeo.ttc"),
    # Linux
    Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"),
    Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
]


def _find_korean_font() -> Optional[Path]:
    env_path = os.getenv("F5_KOREAN_FONT_PATH")
    if env_path:
        p = Path(env_path)
        if p.exists():
            return p
    for p in _KOREAN_FONT_CANDIDATES:
        if p.exists():
            return p
    return None


# 폰트는 프로세스 전역에 한 번만 등록 (중복 등록 방지용 플래그)
_FONTS_REGISTERED: bool = False
_FONT_NAME = "KRFont"


def _ensure_fonts_registered():
    """reportlab 에 한글 TTF 를 전역 등록 (한 번만)."""
    global _FONTS_REGISTERED
    if _FONTS_REGISTERED:
        return

    font_path = _find_korean_font()
    if not font_path:
        raise RuntimeError(
            "한글 폰트를 찾을 수 없습니다. "
            "backend/.env 에 F5_KOREAN_FONT_PATH 를 설정하거나, "
            "시스템에 NanumGothic/Malgun Gothic 등의 한글 TTF 가 설치되어 있어야 합니다."
        )

    pdfmetrics.registerFont(TTFont(_FONT_NAME, str(font_path)))
    _FONTS_REGISTERED = True


def _build_styles():
    """reportlab 스타일 정의. 한글 폰트 전부 적용."""
    styles = getSampleStyleSheet()

    # 제목 (큰 볼드)
    styles.add(ParagraphStyle(
        name="KRTitle",
        fontName=_FONT_NAME,
        fontSize=18,
        leading=22,
        alignment=TA_CENTER,
        textColor=HexColor("#141414"),
        spaceAfter=4,
    ))

    # 부제
    styles.add(ParagraphStyle(
        name="KRSubtitle",
        fontName=_FONT_NAME,
        fontSize=10,
        leading=13,
        alignment=TA_CENTER,
        textColor=HexColor("#787878"),
        spaceAfter=14,
    ))

    # 섹션 헤딩
    styles.add(ParagraphStyle(
        name="KRHeading",
        fontName=_FONT_NAME,
        fontSize=13,
        leading=17,
        textColor=HexColor("#1E4078"),
        spaceBefore=10,
        spaceAfter=5,
    ))

    # 본문
    styles.add(ParagraphStyle(
        name="KRBody",
        fontName=_FONT_NAME,
        fontSize=10,
        leading=14,
        textColor=HexColor("#141414"),
        spaceAfter=3,
    ))

    # 작은 본문 (항목 내부 설명)
    styles.add(ParagraphStyle(
        name="KRSmall",
        fontName=_FONT_NAME,
        fontSize=9,
        leading=12,
        textColor=HexColor("#1E1E1E"),
        leftIndent=12,
        spaceAfter=2,
    ))

    # 항목 타이틀 (bold 느낌)
    styles.add(ParagraphStyle(
        name="KRItemTitle",
        fontName=_FONT_NAME,
        fontSize=11,
        leading=15,
        textColor=HexColor("#282828"),
        spaceBefore=6,
        spaceAfter=2,
    ))

    return styles


def _escape_for_paragraph(text: str) -> str:
    """Paragraph 는 XML 기반이라 <, >, & 를 escape 필요."""
    text = text or ""
    text = text.replace("&", "&amp;")
    text = text.replace("<", "&lt;")
    text = text.replace(">", "&gt;")
    # 줄바꿈은 <br/> 로 변환
    text = text.replace("\n", "<br/>")
    return text


def generate_pdf(step_record: Dict[str, Any]) -> Tuple[bytes, str]:
    """
    F5 리포트를 PDF 바이트로 생성 (reportlab).
    """
    _ensure_fonts_registered()
    data = extract_report_data(step_record)
    styles = _build_styles()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="한글표시사항 검토내역서",
    )

    # reportlab 은 "flowables" 리스트를 build() 에 전달하는 방식
    story = []

    # ── 제목 ──
    story.append(Paragraph("한글표시사항 검토내역서", styles["KRTitle"]))
    story.append(Paragraph("SAMC 수입식품 검역 AI 플랫폼", styles["KRSubtitle"]))

    # ── 1. 기본 정보 ──
    story.append(Paragraph("■ 기본 정보", styles["KRHeading"]))
    info_rows = [
        ["제품명", data["product_name"]],
        ["확정자", data["confirmed_by"] or "(미확정)"],
        ["확정일시", str(data["updated_at"] or "-")],
    ]
    info_table_data = [
        [Paragraph(_escape_for_paragraph(r[0]), styles["KRBody"]),
         Paragraph(_escape_for_paragraph(r[1]), styles["KRBody"])]
        for r in info_rows
    ]
    info_tbl = Table(info_table_data, colWidths=[35 * mm, 120 * mm])
    info_tbl.setStyle(TableStyle([
        ("BOX", (0, 0), (-1, -1), 0.5, HexColor("#C8C8C8")),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, HexColor("#DCDCDC")),
        ("BACKGROUND", (0, 0), (0, -1), HexColor("#F5F5F5")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(info_tbl)

    # ── 2. 교차검증 종합 ──
    story.append(Paragraph("■ 교차검증 종합", styles["KRHeading"]))
    stats = data["stats"]
    summary = (
        f"검토항목 {stats['total']}개  /  "
        f"법령 부적합 {stats['fail']}건  /  "
        f"확인필요 {stats['unclear']}건  /  "
        f"1·2차 불일치 {stats['disagree']}건  /  "
        f"추가 오류 {stats['error']}건"
    )
    story.append(Paragraph(_escape_for_paragraph(summary), styles["KRBody"]))

    # ── 3. 항목별 교차검증 결과 ──
    story.append(Paragraph("■ 항목별 교차검증 결과", styles["KRHeading"]))
    if data["items"]:
        for item in data["items"]:
            field = to_safe_string(item.get("field")) or "-"
            status_kr = STATUS_LABELS.get(item.get("status", ""), item.get("status", "")) or "-"
            v2 = data["v2_by_field"].get(item.get("field")) or {}
            cross_kr = CROSS_LABELS.get(
                v2.get("cross_result", ""), v2.get("cross_result", "") or "-"
            )
            note = to_safe_string(item.get("note"))
            law_ref = to_safe_string(item.get("law_ref"))
            ai_note = to_safe_string(v2.get("ai_note"))

            # 항목 헤더
            header_text = (
                f"<b>■ {_escape_for_paragraph(field)}</b>   "
                f"[1차: {_escape_for_paragraph(status_kr)}]   "
                f"[AI: {_escape_for_paragraph(cross_kr)}]"
            )
            item_block = [Paragraph(header_text, styles["KRItemTitle"])]

            if note:
                item_block.append(
                    Paragraph(
                        f"· <b>검토 의견</b>: {_escape_for_paragraph(note)}",
                        styles["KRSmall"],
                    )
                )
            if ai_note:
                item_block.append(
                    Paragraph(
                        f"· <b>AI 교차검증</b>: {_escape_for_paragraph(ai_note)}",
                        styles["KRSmall"],
                    )
                )
            if law_ref:
                item_block.append(
                    Paragraph(
                        f"· <b>법령 근거</b>: {_escape_for_paragraph(law_ref)}",
                        styles["KRSmall"],
                    )
                )

            # 각 항목을 하나의 덩어리로 - 페이지가 쪼개지면 같이 넘김
            story.append(KeepTogether(item_block))
            story.append(Spacer(1, 3 * mm))
    else:
        story.append(Paragraph("(검토 항목 없음)", styles["KRBody"]))

    # ── 4. AI 추가 발견 이슈 ──
    if data["additional_issues"]:
        story.append(Paragraph("■ AI 추가 발견 이슈", styles["KRHeading"]))
        for issue in data["additional_issues"]:
            sev_kr = SEVERITY_LABELS.get(
                issue.get("severity", ""), issue.get("severity", "")
            )
            field = to_safe_string(issue.get("field"))
            text = to_safe_string(issue.get("issue"))
            line = f"[{_escape_for_paragraph(sev_kr)}] {_escape_for_paragraph(field)} - {_escape_for_paragraph(text)}"
            story.append(Paragraph(line, styles["KRBody"]))

    # ── 5. 최종 한글표시사항 시안 ──
    story.append(Paragraph("■ 최종 한글표시사항 시안", styles["KRHeading"]))
    if data["draft"]:
        draft_table_data = []
        for key, val in data["draft"].items():
            label = DRAFT_LABELS.get(key, key)
            draft_table_data.append([
                Paragraph(_escape_for_paragraph(label), styles["KRBody"]),
                Paragraph(_escape_for_paragraph(to_safe_string(val)) or "-", styles["KRBody"]),
            ])

        draft_tbl = Table(draft_table_data, colWidths=[38 * mm, 117 * mm])
        draft_tbl.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.5, HexColor("#C8C8C8")),
            ("INNERGRID", (0, 0), (-1, -1), 0.3, HexColor("#DCDCDC")),
            ("BACKGROUND", (0, 0), (0, -1), HexColor("#F5F5F5")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(draft_tbl)
    else:
        story.append(Paragraph("(시안 데이터 없음)", styles["KRBody"]))

    # 페이지 푸터는 onPage 콜백으로 추가
    def _on_page(canvas, doc_):
        canvas.saveState()
        canvas.setFont(_FONT_NAME, 8)
        canvas.setFillColor(HexColor("#969696"))
        # 상단 타이틀 (2페이지부터)
        if canvas.getPageNumber() > 1:
            canvas.drawRightString(
                A4[0] - 18 * mm,
                A4[1] - 10 * mm,
                "SAMC 한글표시사항 검토내역서",
            )
        # 하단 페이지 번호
        canvas.drawCentredString(
            A4[0] / 2,
            10 * mm,
            f"Page {canvas.getPageNumber()}",
        )
        canvas.restoreState()

    doc.build(story, onFirstPage=_on_page, onLaterPages=_on_page)

    pdf_bytes = buf.getvalue()
    buf.close()

    filename = make_filename(data["product_name"], data["updated_at"], "pdf")
    return pdf_bytes, filename